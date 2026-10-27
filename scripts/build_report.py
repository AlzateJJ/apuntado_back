"""Genera el Word editable del Entregable 2 a partir del diseño y sus evidencias."""
from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT.parent
MODEL=ROOT/'docs/models';URL='https://github.com/AlzateJJ/apuntado_back/tree/entregable-2'
RISKS=json.loads((ROOT/'docs/risk-storming/risks.json').read_text(encoding='utf-8'))
METRICS=json.loads((ROOT/'docs/evidence/complexity.json').read_text(encoding='utf-8'))
d=Document();plan=[];land=False
d.core_properties.title='Apuntado: diseño detallado, arquitectura de información y gobernanza'
d.core_properties.author='Juan José Alzate Rojas'
d.core_properties.subject='Entregable 2 — Fundamentos de Arquitectura de Software'
for name,size in [('Normal',11),('Title',30),('Subtitle',17),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
    s=d.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.05
    if name.startswith('Heading'):s.font.bold=True;s.paragraph_format.space_before=Pt(8);s.paragraph_format.keep_with_next=True
def geometry(s,l=False):
    s.orientation=WD_ORIENT.LANDSCAPE if l else WD_ORIENT.PORTRAIT
    s.page_width=Inches(11 if l else 8.5);s.page_height=Inches(8.5 if l else 11)
    s.top_margin=Inches(.55);s.bottom_margin=Inches(.55);s.left_margin=Inches(.65);s.right_margin=Inches(.65)
    s.header_distance=Inches(.22);s.footer_distance=Inches(.22)
geometry(d.sections[0])
header=d.sections[0].header.paragraphs[0];header.text='APUNTADO  |  Diseño detallado, información y gobernanza';header.runs[0].font.size=Pt(8)
footer=d.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def page(title,l=False,level=2):
    global land
    if plan:
        if l!=land:geometry(d.add_section(WD_SECTION_START.NEW_PAGE),l)
        else:d.add_page_break()
    land=l;plan.append({'page':len(plan)+1,'title':title,'landscape':l})
    if title:d.add_heading(title,level)
def p(text,bold=False):
    para=d.add_paragraph();r=para.add_run(text);r.bold=bold;return para
def h(text):d.add_heading(text,3)
def table(headers,rows,widths=None,size=10):
    widths=widths or [7.2/len(headers)]*len(headers)
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i]);pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5ECF1' if ri==0 else 'FFFFFF');pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(4);para.paragraph_format.line_spacing=1.0
                for r in para.runs:r.font.size=Pt(size);r.bold=ri==0
    d.add_paragraph().paragraph_format.space_after=Pt(1)
def fig(name,maxh=6.35):
    path=MODEL/(name+'.png');w,hg=Image.open(path).size;maxw=9.7 if land else 7.2;fw=min(maxw,maxh*w/hg)
    para=d.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=para.add_run();shape=run.add_picture(str(path),width=Inches(fw));shape._inline.docPr.set('descr',name.replace('-',' '))
def code(text):
    para=d.add_paragraph();para.paragraph_format.space_after=Pt(6);para.paragraph_format.line_spacing=1.0
    r=para.add_run(text);r.font.name='Consolas';r.font.size=Pt(9)
def link(text,url):
    para=d.add_paragraph();rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    a=OxmlElement('w:hyperlink');a.set(qn('r:id'),rel);r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=text;r.append(t);a.append(r);para._p.append(a)

page('',False)
d.add_paragraph('Apuntado',style='Title')
d.add_paragraph('Diseño detallado, arquitectura de información y gobernanza',style='Subtitle')
p('Segundo entregable del proyecto',True)
p('Juan José Alzate Rojas\nEstudiante de prácticas\nFundamentos de Arquitectura de Software\nUniversidad Nacional de Colombia')
p('El sistema ofrece Apuntado con rondas acumulativas y una modalidad Clásica de partida única. Este informe concreta su arquitectura, el modelo relacional, las decisiones de diseño, los riesgos y las reglas ejecutables de gobernanza.')
p('Se selecciona un monolito modular con estrategias versionadas y persistencia PostgreSQL. La base de código adjunta implementa y prueba el cierre de ronda y sus límites arquitectónicos; la integración de la aplicación multijugador completa continúa en el MVP.')
h('Repositorio del entregable')
link(URL,URL)
p('La rama entregable-2 contiene los modelos editables, ADR-01 a ADR-05, esquema SQL, pruebas de arquitectura y pipeline. Los comandos de reproducción se incluyen al final del informe.')
table(['Contenido','Página'],[
('1. Selección y justificación del estilo arquitectónico','2'),('2. Arquitectura de información y topología de datos','6'),
('3. Registro formal de decisiones arquitectónicas','13'),('4. Análisis de riesgo arquitectónico y Risk Storming','16'),
('5. Gobernanza y funciones de idoneidad','20')],[6.4,.8])

page('1. Selección y Justificación del Estilo Arquitectónico',level=1)
h('1.1 Evaluación Comparativa de Estilos')
p('El Top 3 del primer entregable es Configurability, Testability y Responsiveness. La extensibilidad complementa la configuración: cambiar un parámetro admitido no equivale a crear una familia nueva de reglas. La integridad y la privacidad restringen cualquier alternativa.')
table(['Criterio','Microkernel / plug-ins','Monolito modular'],[
('Configurabilidad','Un registro selecciona extensiones. Requiere contratos y validación de compatibilidad.','Perfiles versionados seleccionan estrategias. Las reglas de cada partida permanecen fijas.'),
('Facilidad de pruebas','Permite probar cada plug-in; también exige pruebas de carga, compatibilidad y aislamiento.','Evalúa reglas puras sin red ni base de datos. La coordinación usa puertos sustituibles.'),
('Capacidad de respuesta','Un plug-in local puede ser rápido; aislarlo en otro proceso introduce comunicación.','Las llamadas internas son locales; el cierre usa una transacción. Debe controlarse la CPU compartida.'),
('Extensión y despliegue','Adecuado para extensiones con ciclo de vida independiente. Carga dinámica y seguridad aumentan el esfuerzo.','Una nueva estrategia se registra y se despliega con el conjunto. El núcleo permanece estable.'),
('Costo individual','Más mecanismos para versionar, cargar y retirar extensiones de los que exige el caso.','Menos infraestructura y una frontera transaccional local, a cambio de disciplina de módulos.')],[1.35,2.9,2.95],10)
p('Se selecciona el monolito modular. Las dos modalidades pertenecen al mismo dominio, comparten turnos y cartas, y no necesitan extensiones instalables por terceros. Strategy y un registro de implementaciones satisfacen la variabilidad actual. La elección reduce coordinación operativa, pero no garantiza tiempos de respuesta sin medirlos.')
p('Se reconsiderará Microkernel si aparecen extensiones de terceros o despliegues independientes. WebSockets es un mecanismo de transporte; su presencia no convierte el sistema en una arquitectura dirigida por eventos.')

page('1.2 Descripción del Estilo Seleccionado y Topología')
p('Un servidor autoritativo contiene los módulos de sala, coordinación, reglas, consulta y sincronización. Persistencia y transporte implementan contratos del sistema. El navegador envía intenciones; el servidor decide validez, turno y puntuación. La topología inicial es un despliegue del backend y una base PostgreSQL central.')
table(['Aspecto','Clásica / partida única','Apuntado'],[
('Duración','Una ronda. El primer cierre válido finaliza la partida y conduce a la vista principal.','Varias rondas con puntuación acumulada.'),
('Perfil inicial','classic@1: cierre 4-3-3, conforme a la modalidad Clásica de la guía.','apuntado@1: 4-3-3, 5-5, diez pintas y toque.'),
('Puntuación','Sin movimientos de puntaje ni eliminación acumulativa.','+10 por combinaciones; +25 por pintas; toque sin premio. Los sobrantes se descuentan.'),
('Fin','Ganador del cierre; no se crea otra ronda.','Victoria en +51 o más, o único activo. Eliminación en -101 o menos.'),
('Reglas compartidas','Dos mazos, 104 cartas, sin jokers, diez cartas al evaluar, hasta seis participantes.','Identidad física, grupos y escaleras; as bajo o alto y valor de 10 puntos.')],[1.15,3.0,3.05],10)
p('El perfil Clásica v1 adopta 4-3-3 como alcance inicial de diseño; los otros cierres se habilitan en Apuntado. La selección y la versión se fijan al crear la partida. Una ampliación de los cierres de Clásica se publicaría como otro perfil, sin cambiar partidas iniciadas.')
h('Empaquetado y extensión')
p('design/src contiene contracts, domain, application, variants, persistence, synchronization, ui y bootstrap. Las variantes solo conocen contratos y utilidades puras. El punto de composición bootstrap/catalog.ts importa sus implementaciones; CloseRound recibe el catálogo por constructor.')
p('Una variante nueva implementa Variant, incorpora sus pruebas y se registra con otra clave o versión. Cambian el registro y el paquete de reglas; no cambia el coordinador. El despliegue sigue siendo conjunto. Esta distinción evita confundir un patrón Strategy con un microkernel dinámico.')

page('1.3 Diagrama UML de Componentes Detallado',True)
fig('uml-componentes',6.35)

page('1.3 Contratos y correspondencia con el código')
table(['Componente','Contrato y responsabilidad','Alcance ejecutable'],[
('C1 Gestionar sala','IConfigurarSala: selecciona perfil, valida participantes y prepara rondas. Requiere repositorio y catálogo.','Modelado; flujo completo de sala pendiente de integración.'),
('C2 Conducir partida','IComando; coordina identidad, turno, fase, versión y cierre atómico mediante RoundStore.','application/closeRound.ts.'),
('C3 Evaluar jugadas','Variant utiliza utilidades de grupos, escaleras, particiones y sobrantes.','domain/ y variants/.'),
('C4 Liquidar ronda','La estrategia calcula cambios y fin de partida sin persistir.','settle de cada variante; resultado tipado Settlement.'),
('C5 Consultar estado','IConsulta: proyecta solo la mano autorizada antes del cierre. El lector SQL se integra por puerto.','application/view.ts sobre un estado recibido.'),
('C6 Persistencia','RoundStore; confirma estado, acta, movimientos, operación y outbox.','persistence/postgres.ts y adaptador de prueba en memoria.'),
('C7 Sincronización','EventStore y AuthorizedPublisher; entrega posterior al commit.','dispatchPending y prueba con publicador sustituible.'),
('Interfaz y transporte','IGameGateway; identidad proveniente de sesión WebSocket y navegación por resultado.','ui/navigation.ts. Servidor WebSocket y pantalla completa pendientes.')],[1.25,3.25,2.7],9.6)
p('El diagrama representa el diseño completo. El código verifica una sección funcional del cierre, no un reemplazo terminado del prototipo. C3 y C4 comparten el paquete de reglas por variante, pero conservan funciones de evaluación y liquidación separadas del coordinador.')
p('IComando corresponde a CloseRound.execute; Variant es el contrato de extensibilidad; RoundStore define la unidad de trabajo. Los nombres ISalaStore e ILeerPartida expresan puertos previstos para completar sala y consultas. La dirección de una llamada no obliga a que el núcleo importe un adaptador concreto.')

page('2. Arquitectura de Información y Topología de Datos',level=1)
h('2.1 Justificación de la Topología')
p('Se adopta una topología relacional centralizada en PostgreSQL 18. El servidor concentra la escritura y coordina cada transición de una partida. Los clientes nunca se conectan directamente a la base. Cada partida puede bloquearse sin bloquear deliberadamente las demás.')
table(['Alternativa','Ventaja para el caso','Costo o límite'],[
('SQL centralizado','PK/FK, consultas de historial y transacción local para acta, puntos y evento.','Administración del servicio, respaldos y posible contención por partida.'),
('SQLite embebido','Instalación sencilla para un proceso y trabajo local.','Concurrencia de escritores más limitada y dificultad al compartir estado entre varias instancias.'),
('Documentos agregados','Una partida incrustada puede actualizarse de forma atómica.','Historial creciente, información duplicada y restricciones de pertenencia menos explícitas.')],[1.35,2.9,2.95],10)
p('La consistencia del marcador pesa más que una instalación mínima: un cierre parcial cambia el resultado del juego. Las transacciones SQL y las claves compuestas permiten expresar pertenencia a partida y unicidad de liquidación. La centralización conserva simplicidad frente a una coordinación distribuida, aunque no aporta alta disponibilidad por sí sola.')
h('2.2 Componente de Información — Opción A Relacional SQL')
p('El modelo autoritativo utiliza once tablas y una vista derivada. Las dos vistas Barker siguientes forman un único modelo. Se repiten los identificadores de algunas entidades para mantener legibilidad; no representan tablas adicionales.')
p('La identidad de una carta es (game_id, card_no). La combinación (game_id, deck, suit, rank) también es única. Dos ases del mismo palo en mazos distintos son copias válidas; repetir una misma copia en dos posiciones de la misma ronda se impide con la PK de card_position.')
p('processed_operation.response contiene únicamente la respuesta serializada para repetir una operación. Es una proyección de transporte subordinada al modelo relacional; no se propone una segunda base documental ni una alternativa NoSQL paralela.')

page('2.2 Modelo E-R Barker — mesa y rondas',True);fig('er-barker-mesa',6.35)
page('2.2 Modelo E-R Barker — liquidación e historial',True);fig('er-barker-liquidacion',6.35)

page('2.2 Normalización y restricciones del esquema')
table(['Relación','Clave y dependencia funcional principal'],[
('rule_profile','profile_id → modo, versión y parámetros. (mode, version) es clave candidata.'),
('player / game','player_id → nombre. game_id → perfil, estado, versión y creación.'),
('participation','(game_id, player_id) → seat. (game_id, seat) también identifica una participación.'),
('round','(game_id, round_no) → estado, fase y jugador del turno.'),
('card_copy','(game_id, card_no) → mazo, palo y rango. Segunda clave: partida + mazo + palo + rango.'),
('card_position','(game_id, round_no, card_no) → zona, dueño y orden. La posición dentro de una zona también es única.'),
('round_result','(game_id, round_no) → declarante, cierre, ganador, siguiente acción y versión confirmada.'),
('score_movement','(game_id, round_no, player_id) → delta. No se almacena nombre ni total acumulado.'),
('processed_operation','(game_id, operation_id) → actor, fingerprint y respuesta.'),
('outbox','event_id → partida, versión, tipo y fechas. (game_id, version, event_type) es clave candidata.')],[1.6,5.6],9.6)
p('Primera forma normal: cada columna representa un valor del contrato; los movimientos y las ubicaciones son filas, no listas incrustadas. Segunda: los atributos no clave dependen de toda la clave compuesta. Tercera: nombres, perfiles y cartas no se duplican en los movimientos; no hay dependencias transitivas de atributos no clave en las tablas autoritativas.')
p('current_score calcula SUM(delta) por participación, con cero inicial. Los valores antes y después de una liquidación se calculan al construir su resultado. Así no se mantienen tres copias independientes del mismo acumulado. La respuesta JSON conserva la representación emitida para idempotencia, no se consulta como fuente del marcador.')
p('El SQL exige una sola ronda abierta por partida, un acta por ronda, un movimiento por jugador y ronda, un asiento por participante y una ubicación por copia. Perfil y asociación de perfil a partida son inmutables. La creación de 104 copias, el mínimo de dos jugadores y la preparación de la ronda se validan en la operación de inicio; las claves por sí solas no garantizan esos mínimos.')

page('2.3 Diagrama UML de Clases de Entidad y Mapeo — mesa',True);fig('uml-entidades-mesa',6.35)
page('2.3 Diagrama UML de Clases de Entidad y Mapeo — liquidación',True);fig('uml-entidades-liquidacion',6.35)

page('2.4 Estrategia de Transaccionalidad y Consistencia')
p('CloseRound.execute recibe partida, ronda, identidad autenticada y comando. El adaptador abre una transacción READ COMMITTED y toma SELECT FOR UPDATE sobre la fila game. Todos los futuros comandos que modifiquen esa partida deben seguir la misma disciplina de bloqueo.')
table(['Paso','Operación e invariante'],[
('1. Bloquear y leer','Bloquear la partida, cargar la ronda y obtener manos y acumulados consistentes. Bloqueos siempre en el mismo orden.'),
('2. Autorizar','Comprobar membresía y turno. Consultar operationId antes de rechazar la versión de un reintento ya confirmado.'),
('3. Deduplicar','Si existe, comparar actor y fingerprint; devolver la respuesta guardada. Reutilización con otra acción: conflicto.'),
('4. Validar','Exigir ronda abierta, fase de decisión y versión esperada. Evaluar la estrategia fijada en el perfil.'),
('5. Liquidar','Conservar el descuento que permitió aceptar el toque. Calcular eliminaciones, ganador y siguiente acción.'),
('6. Persistir','Insertar acta y movimientos; cerrar ronda; actualizar versión y estado; guardar operación y evento pendiente.'),
('7. Confirmar y notificar','COMMIT antes de enviar. El despachador procesa la outbox y marca entrega después de publicar.')],[1.3,5.9],10)
p('Si falla cualquier escritura, ROLLBACK conserva el estado previo. Si falla la red después del commit, la outbox permite reintentar. La entrega puede repetirse: eventId sirve para deduplicar y la versión impide retroceder el estado del cliente. No se promete entrega exactamente una vez.')
p('Ante dos cierres, el segundo espera el bloqueo y observa la nueva versión. Un reintento de la misma operación obtiene la misma acta; una operación distinta encuentra la ronda cerrada. Las PK y restricciones únicas actúan como una segunda barrera.')
p('La prueba PostgreSQL induce una falla al insertar la outbox, después de escribir acta y puntos, y comprueba la reversión. También dispara dos operaciones concurrentes mediante conexiones distintas. Las pruebas de memoria sirven para el dominio; la prueba de PostgreSQL verifica el mecanismo real de bloqueo.')

def adr_page(number,file):
    title='3. Registro Formal de Decisiones Arquitectónicas (ADRs)' if number==3 else f'3. ADR-0{number}'
    page(title,level=1 if number==3 else 2)
    text=(ROOT/'docs/adrs'/file).read_text(encoding='utf-8')
    for block in text.split('\n\n'):
        block=block.strip()
        if not block:continue
        if block.startswith('#'):
            lines=block.splitlines()
            h(lines[0].lstrip('# '))
            if len(lines)>1:p('\n'.join(lines[1:]))
        else:p(block)
adr_page(3,'ADR-03-estilo-extensibilidad.md')
adr_page(4,'ADR-04-persistencia.md')
adr_page(5,'ADR-05-comunicacion.md')

page('4. Análisis de Riesgo Arquitectónico y Risk Storming',level=1)
h('4.1 Matriz de Riesgo 3x3')
p('La prioridad se calcula como impacto × probabilidad. Son valoraciones cualitativas del escenario de diseño; no son frecuencias observadas ni porcentajes estadísticos. Una misma escala se aplica antes y después de proponer mitigaciones.')
table(['Valor','Probabilidad','Impacto'],[
('1','Poco probable; exige una condición excepcional o varias fallas simultáneas.','Efecto local y recuperable, sin cambiar puntuación ni exponer información.'),
('2','Plausible durante una sesión o un cambio de código.','Interrumpe una ronda o dificulta seriamente mantenimiento y respuesta.'),
('3','Esperable por reintentos, concurrencia o cambios frecuentes sin controles.','Altera ganador o marcador, expone manos o compromete la integridad de la partida.')],[.5,3.2,3.5],9.5)
fig('matriz-riesgo',3.6)
p('Los identificadores R1–R9 se mantienen en el registro, las mitigaciones y sus pruebas. Las zonas rojas requieren un cambio arquitectónico o una restricción verificable antes de liberar la funcionalidad afectada.')

page('4.2 Reporte de Risk Storming')
p('Este entregable se desarrolla individualmente debido a mi condición de estudiante de prácticas y a las limitaciones para asistir a la mayoría de las clases, situación que comuniqué previamente al profesor Fernan por correo. Por ello, el ejercicio de Risk Storming recoge únicamente mi perspectiva.')
h('Fase 1 Identificación individual')
p('El análisis recorre las fronteras cliente-servidor, coordinador-reglas, aplicación-persistencia y commit-notificación. Para cada frontera se consideran entradas inválidas, repetición, concurrencia, falla parcial y exposición de información. Se revisan también las diferencias entre modalidades y el mantenimiento de los validadores.')
table(['Grupo de escenarios','Riesgos identificados'],[
('Integridad y concurrencia','R1 doble liquidación; R4 mezcla de reglas; R7 comandos no autorizados u obsoletos.'),
('Privacidad y comunicación','R2 exposición de manos; R3 estado confirmado sin notificación recuperable.'),
('Evolución y capacidad','R5 ciclos; R6 complejidad; R8 evaluación lenta; R9 puntos ciegos de revisión individual.')],[2.1,5.1],10)
h('Fase 2 Consenso — adaptación individual')
p('No hubo consenso grupal ni votación de otros integrantes. Esta fase se adapta a una segunda revisión individual: se agrupan duplicados, se aplica la misma escala y se contrasta cada valoración con su consecuencia concreta y los controles existentes. La justificación queda registrada por riesgo para permitir su revisión posterior.')
p('Se priorizan los escenarios que afectan puntuación, privacidad o autoridad del servidor. R8 permanece medio porque la carga real aún no está medida. La ausencia de revisión por pares se registra como R9; una comprobación automática no reemplaza esa perspectiva adicional.')
h('Fase 3 Mitigación')
p('Las respuestas se incorporan al diseño: bloqueo e idempotencia, perfiles inmutables, proyección privada, outbox y restricciones de dependencias. Cada mecanismo se vincula con una prueba o con una verificación de integración pendiente. El riesgo residual es una estimación condicionada a mantener esos controles, no una garantía de eliminación del riesgo.')
p('Responsable de la valoración, del registro y del seguimiento: Juan José Alzate Rojas. La base de riesgos editable se conserva en docs/risk-storming/risks.json.')

def risk_pages(ids,title):
    page(title)
    for rid in ids:
        r=next(x for x in RISKS if x['id']==rid)
        h(f"{rid} {r['risk']}")
        p(f"Inicial: P={r['p']}, I={r['i']}, nivel={r['p']*r['i']}. {r['why']}")
        p('Mitigación: '+r['mitigation'])
        p(f"Verificación: {r['check']} Residual estimado: P={r['rp']}, I={r['ri']}, nivel={r['rp']*r['ri']}.")
risk_pages(['R1','R2','R3','R4'],'4.3 Plan de Mitigación de Riesgos Críticos')
risk_pages(['R5','R6','R7'],'4.3 Mitigación estructural y seguimiento')
h('R8 y R9 Riesgos que requieren seguimiento')
p('R8: P=2, I=2, nivel 4. Se mantiene residual 4 hasta medir CPU y latencia bajo carga. R9: P=3, I=2, nivel 6. El recorrido sistemático de escenarios y la revisión de resultados reducen la estimación a P=2, I=2, nivel 4; permanece la limitación de una sola perspectiva.')
p('Se reabre el análisis cuando cambia una variante, una transacción, un contrato o la topología. Un resultado rojo sin mitigación bloquea la liberación de la funcionalidad afectada. La aceptación de un riesgo residual debe registrar motivo, responsable y condición de revisión.')

page('5. Gobernanza y Funciones de Idoneidad (Fitness Functions)',level=1)
h('5.1 Reglas de Gobernanza')
table(['Regla','Comprobación y criterio','Respuesta al incumplimiento'],[
('G1 Aislamiento','UI solo usa contratos; no importa persistencia ni drivers SQL. FF-01, FF-04 y FF-05.','Reubicar acceso detrás de un puerto. La verificación falla.'),
('G2 Independencia','Cero ciclos de archivos y cero imports entre classic y apuntado. FF-02 y FF-03.','Extraer contratos o lógica compartida al dominio puro.'),
('G3 Complejidad','CC clásica de McCabe ≤10 por función en domain, variants y application. FF-06.','Descomponer responsabilidades conservando pruebas de comportamiento.'),
('G4 Núcleo estable','Aplicación no importa variantes concretas ni infraestructura. FF-04.','Registrar implementaciones únicamente en bootstrap.'),
('G5 Perfil inmutable','Versión fijada en partida; rechazo de modificación de perfil publicado o de su asociación.','Crear una versión nueva para partidas futuras.'),
('G6 Atomicidad','Acta, movimientos, operación y outbox se confirman o revierten juntos.','Bloquear integración si falla la prueba transaccional.')],[1.3,3.5,2.4],9.7)
p('Juan José Alzate Rojas mantiene decisiones y excepciones. Cada cambio que atraviese un límite requiere actualizar el ADR pertinente y ejecutar el pipeline. Las excepciones deben identificar alcance, razón, responsable y revisión; no se aceptan desactivaciones silenciosas de pruebas o del límite de complejidad.')
p('Estas funciones son automáticas y se ejecutan por cambio en CI. Comprueban propiedades estructurales y algunas invariantes del cierre. La latencia multijugador y la autorización del transporte requieren ensayos adicionales cuando el adaptador WebSocket esté integrado.')

page('5.2 Código Fuente de las Fitness Functions')
p('Se utiliza un solo stack: TypeScript con el ejecutor de pruebas de Node, TSArch y ESLint. El código completo está en tests/architecture/fitness.test.ts. Se usan patrones anclados porque los paths del grafo son relativos al tsconfig; esto evita selectores que no encuentren archivos y aprueben sin comprobar nada.')
code("""const files = () => filesOfProject(CONFIG);

test('UI sin acceso a persistencia', async () => {
  const result = await files()
    .matchingPattern('^design/src/ui/.*')
    .shouldNot().dependOnFiles()
    .matchingPattern('^design/src/persistence/.*').check();
  assert.deepEqual(result, []);
});

test('Cero ciclos', async () => {
  const result = await files()
    .matchingPattern('^design/src/.*')
    .should().beFreeOfCycles().check();
  assert.deepEqual(result, []);
});

test('Variantes independientes', async () => {
  for (const [from, to] of [
    ['classic', 'apuntado'], ['apuntado', 'classic']
  ]) {
    const result = await files()
      .matchingPattern(`^design/src/variants/${from}/.*`)
      .shouldNot().dependOnFiles()
      .matchingPattern(`^design/src/variants/${to}/.*`)
      .check();
    assert.deepEqual(result, []);
  }
});""")
p('La matriz completa permite contracts como base común; domain solo depende de contracts; application no depende de variantes concretas; variants depende de domain y contracts. persistence y synchronization implementan puertos. bootstrap compone las implementaciones.')
p('Además se examinan las dependencias externas prohibidas y se exige que existan los módulos esperados. Las pruebas se aplican a design/src. El directorio src del prototipo histórico se conserva como antecedente y no queda certificado por estos resultados.')

page('5.2 Evidencia ejecutada y controles negativos')
table(['Verificación','Resultado y alcance'],[
('Compilación TypeScript','Sin errores en la base de diseño y las pruebas.'),
('Reglas y aplicación','16 pruebas aprobadas: as, duplicados, cierres, toque, umbrales, permisos, reintentos, rollback y extensión.'),
('Arquitectura','9 pruebas aprobadas, incluidos controles de alcance y casos negativos.'),
('SQL local','1 prueba integral aprobada con PGlite: DDL, claves, restricciones, acumulados y rollback.'),
('PostgreSQL en CI','Prueba de cierre con dos conexiones y falla inducida de outbox aprobada en PostgreSQL 18.'),
('Complejidad',f"{len(METRICS['functions'])} funciones medidas; máximo {METRICS['maximum']}; umbral 10."),
('Reproducción','Versiones fijadas en package-lock.json; Node 24; configuración y comandos incluidos.')],[1.55,5.65],10)
h('Demostración de sensibilidad de las pruebas')
p('Las muestras deliberadamente inválidas se alojan en tests/fixtures/invalid y no forman parte del código de producción. Una UI importa persistencia, dos archivos forman un ciclo y una variante importa la otra. La prueba exige encontrar infracciones en cada muestra.')
code("""const invalid = () => filesOfProject(FIXTURE_CONFIG);
const violations = await invalid()
  .matchingPattern('^ui/.*').shouldNot().dependOnFiles()
  .matchingPattern('^persistence/.*').check();
assert.ok(violations.length > 0);

// Otro control confirma que un import real UI -> contracts
// sí es encontrado por el selector utilizado en la prueba.""")
p('El control de complejidad genera una función con CC 11 y exige un error; otra con CC 10 debe aprobar. Así se verifica el límite inclusivo. La idempotencia también se comprueba después de serializar el resultado como JSON, tal como ocurre en persistencia.')
p('Los archivos verification.json y complexity.json contienen resultados locales y medidas. La ejecución remota se conserva en GitHub Actions. Estas evidencias no equivalen a una prueba de carga con seis navegadores ni a la validación final de la experiencia de usuario.')

page('5.3 Verificación de Complejidad Ciclomática')
p('ESLint analiza la complejidad clásica de McCabe por función. Se aplica a domain, variants y application, incluyendo los validadores y la coordinación del cierre. No se utiliza un promedio por archivo: una sola función que supere 10 hace fallar la comprobación.')
code("""import tseslint from 'typescript-eslint';
export default [{
  files: [
    'design/src/domain/**/*.ts',
    'design/src/variants/**/*.ts',
    'design/src/application/**/*.ts'
  ],
  languageOptions: { parser: tseslint.parser },
  rules: {
    complexity: ['error', { max: 10, variant: 'classic' }]
  }
}];""")
top=sorted(METRICS['functions'],key=lambda f:f['complexity'],reverse=True)[:5]
table(['Función o ubicación','CC medida'],[(x['file'].replace('design/src/','')+':'+str(x['line']),str(x['complexity'])) for x in top],[6.2,1],10)
p('scripts/complexity-report.ts obtiene los valores desde la misma regla de ESLint y conserva archivo y línea. La ejecución de aceptación usa max=10; la extracción de métricas usa max=0 para obtener el valor de todas las funciones y luego comprueba el umbral. No se emplean constantes que simulen una medición.')
h('Integración en el pipeline')
p('GitHub Actions instala desde el lockfile, comprueba tipos, ejecuta dominio, arquitectura y esquema, verifica CC y genera métricas. Después ejecuta la prueba contra PostgreSQL 18 en un servicio efímero. Cualquier proceso con código distinto de cero hace fallar el job; las evidencias se publican incluso si alguna etapa falla.')

page('5.3 Reproducción y alcance de la entrega')
code("""git clone --branch entregable-2 \\
  https://github.com/AlzateJJ/apuntado_back.git
cd apuntado_back
# Usar Node.js 24
npm ci --ignore-scripts
npm run verify""")
p('La prueba de PostgreSQL utiliza ARCH_TEST_DATABASE_URL y una base exclusiva y vacía. En CI el servicio se crea y elimina para cada ejecución. Los comandos locales de verificación no usan la base ni el archivo .env del prototipo.')
table(['Artefacto','Ubicación en el repositorio'],[
('Modelos UML y Barker','docs/models/: drawio, SVG, PDF vectorial, PNG y entities.json.'),
('Decisiones','docs/adrs/ADR-01…ADR-05; antecedente de WebSockets en docs/history/.'),
('Información','database/schema.sql; tipos y relaciones en los modelos de entidades.'),
('Riesgos','docs/risk-storming/risks.json y reporte individual de este documento.'),
('Fitness functions','tests/architecture/fitness.test.ts y eslint.architecture.config.mjs.'),
('Código verificado','design/src/; pruebas de reglas, aplicación, SQL y PostgreSQL en tests/.'),
('Integración continua','Workflow architecture.yml; resultados en docs/evidence/ y GitHub Actions.')],[1.7,5.5],10)
p('La entrega comprueba la estructura modular y el camino de cierre de ronda. Quedan para integrar el MVP: sala y reparto completos, robo y descarte, autenticación del transporte, servidor WebSocket, reconexión entre dispositivos y pantallas de juego. La navegación de partida única ya está expresada y probada mediante el resultado home, cuyo destino es la vista principal.')
p('Se conservan como metas del primer entregable evaluación p95 ≤50 ms y confirmación y difusión p95 ≤500 ms en el escenario de piloto. Las pruebas actuales no constituyen la medición de esas metas. Su futura ejecución debe registrar carga, red, muestra y percentiles por separado.')
p('El informe consolida el diseño, los riesgos y los controles. Los archivos fuente permiten modificar modelos y ejecutar comprobaciones; cualquier cambio posterior debe mantener alineados código, esquema, diagramas y ADR.')

OUT.mkdir(exist_ok=True);dest=OUT/'Apuntado_Entregable_2.docx';d.save(dest)
(ROOT/'docs/evidence/page-plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
print(str(dest));print('Páginas planificadas:',len(plan))
